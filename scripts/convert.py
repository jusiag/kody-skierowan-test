"""Konwersja słownika ICD-9 PL CM z NFZ (plik .xlsx) na zwarty plik JSON dla aplikacji.

Układ pliku NFZ (sprawdzony na wersji 5.90):
  * główne dane: PIERWSZY arkusz, którego nazwa (bez spacji) zaczyna się od "ICD-9",
    dwa wiersze nagłówka, potem kolumny A-H jako cztery pary (kod, nazwa):
    rozdział, podrozdział, kategoria główna, kategoria szczegółowa;
  * w niektórych wierszach kolumna rozdziału jest pusta, a dalsze są wypełnione
    (np. L68, T28), więc puste komórki pomijamy, a nie przerywamy czytania;
  * część kodów (88, 93) jest zapisana jako liczba, więc zawsze zamieniamy na tekst;
  * arkusze "wł_wył..." zawierają uwagi "Obejmuje" / "Nie obejmuje",
    arkusz "Uwagi" zawiera uwagi do kodowania.
"""
import json
import re
import sys

import openpyxl

MIN_CODES = 5000          # poniżej tej liczby plik uznajemy za niepełny
REQUIRED_CODES = ("89.00", "88.713", "87.440")


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def version_key(v):
    return tuple(int(p) for p in re.findall(r"\d+", str(v))) or (0,)


def version_from_name(name):
    """Wyciąga numer wersji (np. 5.90) z nazwy pliku NFZ."""
    m = re.search(r"(?<!\d)(\d{1,2})[._](\d{2})(?!\d)", name)
    return f"{m.group(1)}.{m.group(2)}" if m else None


def convert(path, version):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    main_ws = next(
        (ws for ws in wb.worksheets if ws.title.strip().upper().startswith("ICD-9")), None
    )
    if main_ws is None:
        raise ValueError(
            "W pliku nie ma arkusza zaczynającego się od 'ICD-9'. "
            "To nie wygląda na słownik ICD-9 PL CM z nfz.gov.pl."
        )

    codes = {}
    order = []
    for row in main_ws.iter_rows(min_row=3, values_only=True):
        r = [clean(c) for c in row[:8]] + [""] * 8
        parent = ""
        for lvl, i in enumerate((0, 2, 4, 6)):
            code, name = r[i], r[i + 1]
            if not code:
                continue
            if code not in codes:
                codes[code] = [name, lvl, parent]
                order.append(code)
            elif name and not codes[code][0]:
                codes[code][0] = name
            parent = code

    notes = {}

    def add(code, kind, text):
        text = clean(text)
        if not code or not text or code not in codes:
            return
        lst = notes.setdefault(code, {}).setdefault(kind, [])
        if text not in lst:
            lst.append(text)

    for ws in wb.worksheets:
        t = ws.title.strip().lower()
        if t.startswith("wł_wył") or t.startswith("wl_wyl"):
            for row in ws.iter_rows(min_row=2, values_only=True):
                kind = clean(row[2] if len(row) > 2 else "").lower()
                text = row[3] if len(row) > 3 else None
                if kind.startswith("nie"):
                    add(clean(row[0]), "x", text)
                elif kind.startswith("obejm"):
                    add(clean(row[0]), "i", text)
        elif t.startswith("uwagi"):
            for row in ws.iter_rows(min_row=2, values_only=True):
                add(clean(row[0]), "u", row[2] if len(row) > 2 else None)

    has_children = {v[2] for v in codes.values() if v[2]}
    data = {
        "version": version,
        "count": len(order),
        "codes": [
            [c, codes[c][0], codes[c][1], codes[c][2], 1 if c in has_children else 0]
            for c in order
        ],
        "notes": notes,
    }
    validate(data)
    return data


def validate(data):
    if data["count"] < MIN_CODES:
        raise ValueError(
            f"Plik wygląda na niepełny: tylko {data['count']} kodów (oczekiwano co najmniej {MIN_CODES})."
        )
    have = {c[0] for c in data["codes"]}
    missing = [c for c in REQUIRED_CODES if c not in have]
    if missing:
        raise ValueError("W słowniku brakuje kodów, które muszą tam być: " + ", ".join(missing))
    if any(not c[1] for c in data["codes"] if c[2] >= 2):
        raise ValueError("Część kodów szczegółowych nie ma nazwy, plik jest uszkodzony lub ma inny układ.")


def write_json(data, out):
    with open(out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))


if __name__ == "__main__":
    src, version, out = sys.argv[1:4]
    d = convert(src, version)
    write_json(d, out)
    print(f"{d['count']} kodów, {len(d['notes'])} z uwagami -> {out}")
